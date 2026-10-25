#!/usr/bin/env python3
"""Build the 5-minute report deck for the low-latency enhancement terminal.

Nine 16:9 slides, Chinese, speaker notes with a timing plan. Every number is taken from a file in
this repository:

    board/hdmi_stream_enhance/evidence/final_measurements.json   WNS, latency, offset, stability
    board/dma_sobel/metrics_board.csv                            PS vs PL frame times
    build/hls/enhance/reports/csynth.rpt                         resources / estimated Fmax
    docs/output_offset.md                                        per-stream offset (216.8 -> 218, residual 0-1 px)

Figures are existing board evidence (no synthetic renders):

    board/hdmi_stream_enhance/evidence/presets/brightness_ladder.png
    board/hdmi_stream_enhance/evidence/hdmi_final_{a_plain,b_brighter,f_split}.png
    board/hdmi_stream_enhance/evidence/tuning/source_dark_page.jpg

    python report/slides/make_slides.py
"""

from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches as In, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(HERE, "低延迟视觉增强终端_5分钟汇报.pptx")
TMP = os.path.join(REPO, ".tmp", "slides")

EVID = os.path.join(REPO, "board", "hdmi_stream_enhance", "evidence")
IMG_LADDER = os.path.join(EVID, "presets", "brightness_ladder.png")
IMG_PLAIN = os.path.join(EVID, "hdmi_final_a_plain.png")
IMG_BRIGHT = os.path.join(EVID, "hdmi_final_b_brighter.png")
IMG_SPLIT = os.path.join(EVID, "hdmi_final_f_split.png")
IMG_DARK = os.path.join(REPO, "board", "hdmi_stream_enhance", "evidence", "tuning", "source_dark_page.jpg")

FONT = "Microsoft YaHei"

NAVY = RGBColor(0x0F, 0x22, 0x39)
INK = RGBColor(0x16, 0x22, 0x30)
ACCENT = RGBColor(0x1B, 0x7F, 0xD4)
SKY = RGBColor(0x63, 0xAE, 0xE8)
AMBER = RGBColor(0xD9, 0x83, 0x2B)
GREY = RGBColor(0x55, 0x64, 0x75)
MUTED = RGBColor(0x86, 0x95, 0xA6)
LINE = RGBColor(0xD9, 0xE3, 0xEC)
CARD = RGBColor(0xF4, 0xF8, 0xFC)
DEEP = RGBColor(0x18, 0x33, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
M = 0.72
CW = W - 2 * M

SECTION = "低延迟视觉增强终端  ·  PYNQ-Z2  ·  2026 嵌入式芯片与系统设计竞赛"


# --------------------------------------------------------------------------- primitives


def _fix_fonts(run) -> None:
    """python-pptx only sets the Latin face; Chinese needs a:ea as well."""
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set("typeface", FONT)


def run(p, text, size, bold=False, color=INK, italic=False, spacing=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = FONT
    if spacing is not None:
        r._r.get_or_add_rPr().set("spc", str(spacing))
    _fix_fonts(r)
    return r


def frame(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, pad=(0, 0, 0, 0), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = In(pad[0])
    tf.margin_top = In(pad[1])
    tf.margin_right = In(pad[2])
    tf.margin_bottom = In(pad[3])
    tf.paragraphs[0].alignment = align
    return tf


def line(tf, items, first=False, space_before=0, space_after=0, spacing=1.18, align=PP_ALIGN.LEFT):
    """items: [(text, size, bold, color)] or a single tuple; one paragraph per call."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)
    p.line_spacing = spacing
    p.alignment = align
    if isinstance(items, tuple):
        items = [items]
    for text, size, bold, color in items:
        run(p, text, size, bold, color)
    return p


def box(slide, x, y, w, h, fill=None, outline=None, radius=0.09, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if outline is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = outline
        s.line.width = Pt(0.9)
    s.shadow.inherit = False
    if s.has_text_frame:
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = In(0.12)
        tf.margin_right = In(0.12)
        tf.margin_top = In(0.06)
        tf.margin_bottom = In(0.06)
    return s


def arrow(slide, x, y, w=0.24, h=0.16, color=ACCENT):
    box(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def rule(slide, x, y, w, h=0.008, color=LINE):
    box(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.RECTANGLE)


def heading(slide, num, tag, title):
    box(slide, M, 0.80, 0.075, 0.52, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
    tf = frame(slide, M + 0.22, 0.40, CW, 0.32)
    line(tf, [(f"{num}   ", 11, True, ACCENT), (tag, 11, True, MUTED)], first=True, spacing=1.0)
    tf = frame(slide, M + 0.22, 0.70, CW, 0.70)
    line(tf, (title, 28, True, NAVY), first=True, spacing=1.0)


def footer(slide, page, total=9):
    rule(slide, M, 7.04, CW, 0.008, LINE)
    tf = frame(slide, M, 7.12, 10.0, 0.25)
    line(tf, (SECTION, 9, False, MUTED), first=True, spacing=1.0)
    tf = frame(slide, W - M - 1.2, 7.12, 1.2, 0.25)
    line(tf, (f"{page:02d} / {total:02d}", 9, False, MUTED), first=True, spacing=1.0, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = In(w)
    if h:
        kw["height"] = In(h)
    return slide.shapes.add_picture(path, In(x), In(y), **kw)


def framed_picture(slide, path, x, y, w, color=LINE):
    """Place a picture at a fixed width and give it a hairline frame; returns its height."""
    iw, ih = Image.open(path).size
    h = w * ih / iw
    slide.shapes.add_picture(path, In(x), In(y), width=In(w))
    f = box(slide, x, y, w, h, fill=None, outline=color, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    f.line.width = Pt(0.75)
    return h


def upscale(src, factor=2, name=None):
    """Board evidence is small; resample once so the renderer does not show hard pixel steps."""
    os.makedirs(TMP, exist_ok=True)
    dst = os.path.join(TMP, name or (os.path.splitext(os.path.basename(src))[0] + f"_{factor}x.png"))
    if not os.path.exists(dst):
        im = Image.open(src).convert("RGB")
        im = im.resize((im.width * factor, im.height * factor), Image.LANCZOS)
        im.save(dst, optimize=True)
    return dst


def spacer(slide, y, text, size=12, color=MUTED, w=CW, x=M, h=0.3):
    tf = frame(slide, x, y, w, h)
    line(tf, (text, size, False, color), first=True, spacing=1.0)


# --------------------------------------------------------------------------- slides


def slide_cover(prs):
    s = new(prs)
    box(s, 0, 0, W, H, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    box(s, 0, H - 0.16, W, 0.16, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)

    tf = frame(s, M + 0.18, 0.78, 11.6, 0.3)
    line(tf, ("2026 全国大学生嵌入式芯片与系统设计竞赛  ·  FPGA 创新设计赛道  ·  AMD 自主选题（初级组）",
              11.5, True, SKY), first=True, spacing=1.05)

    tf = frame(s, M + 0.18, 1.95, 11.4, 1.0)
    line(tf, ("低延迟视觉增强终端", 52, True, WHITE), first=True, spacing=1.0)

    tf = frame(s, M + 0.20, 2.95, 11.0, 0.5)
    line(tf, ("零帧缓存  ·  像素边走边算  ·  延迟不到一行，而不是一帧", 21, False, RGBColor(0xBF, 0xD2, 0xE6)),
         first=True, spacing=1.0)

    rule(s, M + 0.20, 3.62, 2.3, 0.035, SKY)

    tf = frame(s, M + 0.20, 3.82, 11.0, 0.4)
    line(tf, ("平台  PYNQ-Z2（Zynq-7000 XC7Z020）      工具  Vivado / Vitis HLS 2026.1      "
              "面向  低照度图像观察", 11.5, False, RGBColor(0x93, 0xAB, 0xC4)), first=True, spacing=1.0)

    stats = [("21 个时钟", "内核流水线延迟（≈ 0.15 µs）"),
             ("0 / 2,764,800", "板上逐位比对的通道差错"),
             ("196.7×", "相对 ARM 纯软件实现")]
    x = M + 0.18
    for big, small in stats:
        b = box(s, x, 5.02, 3.46, 1.16, fill=DEEP)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        line(tf, (big, 25, True, WHITE), first=True, spacing=1.0, align=PP_ALIGN.CENTER)
        line(tf, (small, 10, False, RGBColor(0x9D, 0xB5, 0xCD)), space_before=4, spacing=1.0,
             align=PP_ALIGN.CENTER)
        x += 3.46 + 0.20

    tf = frame(s, M + 0.20, 6.60, 11.4, 0.3)
    line(tf, ("汇报 5 分钟  ｜  2026.10  ｜  证据与复现：README.md · report/design_report.md",
              9.5, False, RGBColor(0x74, 0x8A, 0xA2)), first=True, spacing=1.0)

    notes(s, "【0:00–0:20 开场】\n"
             "各位老师好，我的作品是低延迟视觉增强终端。一句话说清它和别的方案的区别："
             "画面不做帧缓存，像素从接口进来就直接算、直接送出去，所以延迟不到一行，而不是一帧。\n"
             "这一页的三个数字就是全篇的三个结论：延迟 21 个时钟、逐位比对 0 差错、相对 ARM 快 196.7 倍。")
    return s


def slide_problem(prs):
    s = new(prs)
    heading(s, "01", "问题", "暗光下的画面，卡的不是帧率，是延迟")

    items = [
        ("场景", "夜间巡检与取证、车载夜视、工业与显微弱光成像。"),
        ("难点", "不是分辨率，而是画面偏暗、反差不足 —— 信息在，但读不出来。"),
        ("现状", "通用手机 / PC 通路在 100–300 ms 量级：迟滞、拖影，长时间观察容易疲劳。"),
    ]
    y = 1.86
    for label, body in items:
        tf = frame(s, M + 0.22, y, 6.55, 0.3)
        line(tf, (label, 11, True, ACCENT), first=True, spacing=1.0)
        tf = frame(s, M + 0.22, y + 0.30, 6.55, 0.75)
        line(tf, (body, 15.5, False, INK), first=True, spacing=1.25)
        y += 1.02

    picture(s, IMG_DARK, 7.55, 1.82, w=4.75, h=3.56)
    tf = frame(s, 7.55, 5.44, 4.75, 0.3)
    line(tf, ("暗光书页：肉眼勉强可读，拍下来几乎全黑", 10, False, MUTED), first=True, spacing=1.0)

    band = box(s, M, 5.74, CW, 0.78, fill=CARD)
    box(s, M, 5.74, 0.07, 0.78, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
    tf = band.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.28)
    line(tf, ("核心指标是端到端延迟，不是帧率 —— 所以整条增强链不进帧缓冲。", 16, True, NAVY),
         first=True, spacing=1.0)

    spacer(s, 6.62, "注：100–300 ms 为公开经验量级，非本项目实测。", 9, MUTED)

    footer(s, 2)
    notes(s, "【0:20–1:00 问题】\n"
             "低照度观察的困难不是看不清细节，而是整幅偏暗、反差被压掉。手机和 PC 的通路通常是"
             "采集、处理、显示中间过帧缓冲，量级在 100–300 毫秒，人一动画面就迟滞。\n"
             "这类设备真正要的是延迟，所以我把整个处理放在像素流里做，不做帧缓存。"
             "代价也很明确：放弃了降噪、稳像这类需要多帧的后处理。")
    return s


def slide_arch(prs):
    s = new(prs)
    heading(s, "02", "方案", "零帧缓存：增强链整个塞进像素流里")

    blocks = [
        ("PS 端", "Python / Jupyter\n送帧、改参数", CARD, NAVY, GREY),
        ("DDR + VDMA", "PS 写帧，VDMA 供帧", CARD, NAVY, GREY),
        ("SOF · enhance", "增强核在真实像素流上", ACCENT, WHITE, RGBColor(0xD6, 0xE9, 0xFA)),
        ("unpack · vid_out", "重建视频时序", CARD, NAVY, GREY),
        ("HDMI OUT", "640×480 @60", CARD, NAVY, GREY),
    ]
    xs = [M, 2.82, 5.12, 8.32, 10.92]
    ws = [1.85, 2.05, 2.95, 2.35, 1.69]
    for (title, sub, fill, fg, subf), x, w in zip(blocks, xs, ws):
        b = box(s, x, 1.88, w, 0.98, fill=fill)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        line(tf, (title, 12.5, True, fg), first=True, spacing=1.0, align=PP_ALIGN.CENTER)
        for i, part in enumerate(sub.split("\n")):
            line(tf, (part, 9.5, False, subf if fill == ACCENT else GREY),
                 space_before=2, spacing=1.0, align=PP_ALIGN.CENTER)
    for x in (2.585, 4.895, 8.075, 10.685):
        arrow(s, x, 2.29, 0.22, 0.16)

    ver = box(s, M, 3.08, CW, 0.62, fill=CARD)
    tf = ver.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.22)
    line(tf, [("验证通路（不需要显示器）  ", 12, True, ACCENT),
              ("DDR ↔ AXI-DMA ↔ enhance → 回读 → numpy 黄金模型逐位比对", 12, False, INK)],
         first=True, spacing=1.0)

    cards = [
        ("零帧缓存", "增强链在像素流内一次走完，运算不落帧缓冲；PS 只负责送帧与改参数，不参与任何像素运算。"),
        ("一份源码 · 两条通路", "DMA 通路可回读、可逐位验证；HDMI 通路在真实像素流上出图。两条通路共用同一份 HLS 源码。"),
        ("帧源为什么是 VDMA", "实测这块板 HDMI 输出只能 640×480，而 HDMI 输入又跑不了 640×480：直通无解。"
                            "改由 VDMA 供帧，内核仍在真实像素通路上。"),
    ]
    x = M
    for i, (title, body) in enumerate(cards):
        b = box(s, x, 3.95, 3.82, 2.30, fill=None, outline=LINE)
        tf = b.text_frame
        tf.margin_left = In(0.2)
        tf.margin_right = In(0.18)
        tf.margin_top = In(0.18)
        line(tf, (f"0{i + 1}", 10, True, SKY), first=True, spacing=1.0)
        line(tf, (title, 14.5, True, NAVY), space_before=3, spacing=1.0)
        line(tf, (body, 11.5, False, GREY), space_before=9, spacing=1.28)
        x += 3.82 + 0.215

    spacer(s, 6.48, "位流：enhance_hdmi.bit（WNS +0.206 ns，142.857 MHz，0 失败端点）", 10, MUTED)

    footer(s, 3)
    notes(s, "【1:00–1:45 方案】\n"
             "数据从 PS 写进 DDR，由 VDMA 供帧，然后依次过 SOF 对齐、增强核、像素解包和视频时序重建，最后出 HDMI。"
             "增强核在真实像素通路上，PS 不碰任何像素。\n"
             "下面这条是验证通路：同一份源码换成 DMA 版本，可以整帧回读，跟 numpy 黄金模型逐位比对，不需要显示器。\n"
             "为什么帧源不是 HDMI 输入？实测这块板的 HDMI 输出只支持 640×480，而输入 IP 又跑不了 640×480，"
             "原厂位流也一样，所以直通无解 —— 这是硬件上限，不是没调通。")
    return s


def slide_kernel(prs):
    s = new(prs)
    heading(s, "03", "核心算子", "一条流水线，59 bit 参数全在运行期可调")

    stages = [
        ("01", "局部对比", "clarity", "先算局部均值，再按比例放大细节"),
        ("02", "两段色调曲线", "knee · gain_lo · gain_hi", "拐点以下、以上分别给不同增益"),
        ("03", "二值 · 反色 · 单色", "thresh · invert · mono", "硬阈值切分，可反色、去色"),
        ("04", "分屏", "split_x", "同一帧内左原始、右处理，列计数器即可"),
    ]
    x = M
    for num, title, params, desc in stages:
        b = box(s, x, 1.86, 2.75, 1.36, fill=CARD)
        tf = b.text_frame
        tf.margin_left = In(0.16)
        tf.margin_top = In(0.12)
        line(tf, (num, 9.5, True, SKY), first=True, spacing=1.0)
        line(tf, (title, 14.5, True, NAVY), space_before=1, spacing=1.0)
        line(tf, (params, 9.5, False, ACCENT), space_before=3, spacing=1.0)
        tf = frame(s, x + 0.04, 3.26, 2.67, 0.5)
        line(tf, (desc, 9.8, False, MUTED), first=True, spacing=1.15)
        x += 2.75 + 0.30

    tf = frame(s, M, 3.86, CW, 0.7)
    line(tf, [("全整数定点、无浮点；每像素 1 拍（II = 1）；", 14.5, False, INK),
              ("三块行缓冲轮转，每像素对每块缓冲一读一写 —— 正好卡进双口 BRAM。", 14.5, True, NAVY)],
         first=True, spacing=1.1)

    chips = [("3 BRAM · 2 DSP", "1659 LUT，占 3%"),
             ("194.59 MHz", "7 ns 约束下的估计 Fmax"),
             ("59 bit", "运行期参数，PS 下发"),
             ("10 个", "预设，运行期一键切换")]
    x = M
    for big, small in chips:
        b = box(s, x, 4.62, 2.85, 1.06, fill=None, outline=LINE)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        line(tf, (big, 16, True, ACCENT), first=True, spacing=1.0, align=PP_ALIGN.CENTER)
        line(tf, (small, 10, False, GREY), space_before=3, spacing=1.0, align=PP_ALIGN.CENTER)
        x += 2.85 + 0.155

    b = box(s, M, 5.94, CW, 0.86, fill=CARD)
    box(s, M, 5.94, 0.07, 0.86, fill=AMBER, shape=MSO_SHAPE.RECTANGLE)
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.28)
    line(tf, [("参数每像素从 BRAM 读一次 → 切预设只是一次寄存器写，像素流不中断。", 13.5, True, NAVY),
              ("（这是被 HLS 逼出来的设计，见第 07 页）", 11.5, False, GREY)], first=True, spacing=1.1)

    footer(s, 4)
    notes(s, "【1:45–2:20 算子】\n"
             "算子只有一条流水线，四段：局部对比、两段色调曲线、二值/反色/单色、分屏。"
             "所有参数一共 59 位，PS 在运行期下发，一共 10 个预设。\n"
             "实现上全整数定点，每像素一拍，三块行缓冲轮转，每像素对每块缓冲只读一次写一次，"
             "所以刚好塞进双口 BRAM。资源是 3 块 BRAM、2 个 DSP、1659 个 LUT，其中 LUT 只占 3%。\n"
             "分屏是免费的：用现成的列计数器就能在同一帧里左右分开，这是答辩现场最直观的证据。")
    return s


def slide_effect(prs):
    s = new(prs)
    heading(s, "04", "效果", "同一张暗图：平均亮度 29 → 96 / 100 / 104")

    tf = frame(s, M + 0.22, 1.42, CW, 0.4)
    line(tf, ("pass 是恒等（不处理）；bright / document / night 分别针对提亮、文档与夜间，高光截断都在 0.2% 以下。",
              12.5, False, GREY), first=True, spacing=1.05)

    pic = upscale(IMG_LADDER, 2)
    framed_picture(s, pic, 1.77, 1.80, 9.8)
    tf = frame(s, 1.77, 6.22, 9.8, 0.3)
    line(tf, ("同一张暗图、四个预设的板上实测（采集卡 640×480）；每格上方的数字是该预设下的平均亮度",
              9.8, False, MUTED), first=True, spacing=1.0, align=PP_ALIGN.CENTER)

    tf = frame(s, M, 6.54, CW, 0.36)
    line(tf, [("极性保持是特意调出来的：", 10.5, True, AMBER),
              ("暗页提亮不能把黑白翻过来 —— 旧 reading 预设对比度 +41 → −51，新 document 预设 +158。",
               10.5, False, GREY)], first=True, spacing=1.0)

    footer(s, 5)
    notes(s, "【2:20–2:55 效果】\n"
             "这张图是同一张暗图在四个预设下的实拍，右边三张分别是提亮、文档、夜间。平均亮度从 29 提到 96 到 104，"
             "高光截断都在 0.2% 以下。\n"
             "一个踩过的坑写在这里：最早的阅读预设在暗页上会把黑白翻过来，对比度从 +41 掉到 −51；"
             "后来用可复现的脚本重新搜参数，才有了 document 预设，对比度 +158。\n"
             "也就是说，“提亮”不等于“可读”，极性必须保住。")
    return s


def slide_correctness(prs):
    s = new(prs)
    heading(s, "05", "正确性", "不是“看着像”，是逐位一致")

    cards = [
        ("C 仿真", "9216 拍", "0 误差", "对独立编写的 C++ 黄金模型，两种配置各推 9216 拍，侧带信号完好。"),
        ("板上回读", "0 / 2,764,800", "通道不一致", "AXI-DMA 整帧回环，8 档（验证参数表）逐一比 numpy：每档 0 个通道不一致，15.6 ms/帧。"),
        ("边界策略", "一套口径", "三处同源", "水平复制、垂直用上一帧行 —— 写进头文件、测试台与 Python 参考，所以能要求 0 容差。"),
    ]
    x = M
    for label, big, mid, body in cards:
        b = box(s, x, 1.80, 3.82, 2.78, fill=CARD)
        tf = b.text_frame
        tf.margin_left = In(0.24)
        tf.margin_right = In(0.22)
        tf.margin_top = In(0.24)
        line(tf, (label, 10, True, MUTED), first=True, spacing=1.0)
        line(tf, (big, 25, True, ACCENT), space_before=6, spacing=1.0)
        line(tf, (mid, 12, True, NAVY), space_before=4, spacing=1.0)
        line(tf, (body, 11.5, False, GREY), space_before=9, spacing=1.30)
        x += 3.82 + 0.215

    b = box(s, M, 4.96, CW, 0.90, fill=None, outline=LINE)
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.3)
    line(tf, ("只有 0 容差比对，才配说“结果正确”；“看起来差不多”不算验证，也骗不过边界。",
              15, True, NAVY), first=True, spacing=1.0)

    tf = frame(s, M, 6.04, CW, 0.3)
    line(tf, ("证据路径：src/hls/enhance/tb_enhance*.cpp  ·  board/enhance_dma/  ·  sim/enhance/golden_enhance.py",
              9.5, False, MUTED), first=True, spacing=1.0)

    footer(s, 6)
    notes(s, "【2:55–3:30 正确性】\n"
             "硬件算子最容易停在“看起来对”。我这里做了三层：C 仿真对一份独立写的 C++ 黄金模型，两种配置各 9216 拍，"
             "0 误差；板上把整帧回读，8 档（验证参数表）逐一比 numpy，每档 276 万多个通道，0 个不一致；"
             "边界策略也一样——水平复制、垂直用上一帧行，硬件、测试台、Python 三处同口径。\n"
             "能要求 0 容差，是因为先把边界这件事定义死了，而不是靠事后调容差。\n"
             "（如果有人问 10 个预设为什么这里写 8 档：DMA 逐位验证用的那一版参数表是 8 档，"
             "而且 night/document 之后还调过；HDMI 侧现在运行期是 10 档，但那 10 档没有逐位验证。）")
    return s


def slide_perf(prs):
    s = new(prs)
    heading(s, "06", "性能", "同一块板，ARM 还是 PL")

    tf = frame(s, M + 0.18, 1.62, 5.4, 1.55)
    line(tf, ("196.7×", 86, True, ACCENT), first=True, spacing=0.95)
    tf = frame(s, M + 0.24, 3.10, 5.2, 0.9)
    line(tf, ("同一块 PYNQ-Z2 · 同一张 1280×720 图 · 同一个 Sobel\n两边输出逐像素完全一致",
              12, False, GREY), first=True, spacing=1.3)
    tf = frame(s, M + 0.24, 3.72, 5.2, 0.4)
    line(tf, ("同一份算法描述 · 同一套测试图 · 同一份测量脚本", 10.5, False, MUTED),
         first=True, spacing=1.0)

    # proportional bars: 3108.4 ms vs 15.8 ms
    base_x, base_w = 6.55, 5.55
    tf = frame(s, base_x, 1.72, base_w, 0.3)
    line(tf, ("PS：ARM Cortex-A9 + numpy", 11, True, NAVY), first=True, spacing=1.0)
    box(s, base_x, 2.02, base_w, 0.40, fill=RGBColor(0xC4, 0xD2, 0xDF), radius=0.02)
    tf = frame(s, base_x, 2.02, base_w - 0.14, 0.40, anchor=MSO_ANCHOR.MIDDLE)
    line(tf, ("3108.4 ms / 帧", 12, True, NAVY), first=True, spacing=1.0, align=PP_ALIGN.RIGHT)

    tf = frame(s, base_x, 2.66, base_w, 0.3)
    line(tf, ("PL：Zynq 逻辑 + AXI-DMA 整帧传输", 11, True, ACCENT), first=True, spacing=1.0)
    box(s, base_x, 2.96, 0.055, 0.40, fill=ACCENT, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    tf = frame(s, base_x + 0.2, 2.96, base_w - 0.2, 0.40, anchor=MSO_ANCHOR.MIDDLE)
    line(tf, ("15.8 ms / 帧  ·  63.45 fps", 12, True, ACCENT), first=True, spacing=1.0)

    tf = frame(s, base_x, 3.50, base_w, 0.3)
    line(tf, ("同比例尺下，PL 的条只有 0.03 英寸宽 —— 所以画成一条线", 9.5, False, MUTED),
         first=True, spacing=1.0)

    chips = [("15.6 ms/帧", "增强核整帧（DMA，8 档验证表逐位比对）"),
             ("21 个时钟", "内核流水线延迟 ≈ 0.15 µs，不到一行"),
             ("+0.206 ns", "WNS @142.857 MHz，0 失败端点")]
    x = M
    for big, small in chips:
        b = box(s, x, 4.58, 3.82, 1.20, fill=CARD)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        line(tf, (big, 19, True, NAVY), first=True, spacing=1.0)
        line(tf, (small, 10.5, False, GREY), space_before=4, spacing=1.1)
        x += 3.82 + 0.215

    tf = frame(s, M, 6.08, CW, 0.8)
    line(tf, ("逐行的 659 ms 是软件开销（720 次核启动 + 1440 次 DMA），改整帧传输后降到 15.8 ms。",
              10, False, GREY), first=True, spacing=1.2)
    line(tf, ("主机 numpy 40.7 ms 更快 —— 所以比的是同一块嵌入式板卡上的 ARM 与 PL，而非桌面 CPU。",
              10, False, GREY), spacing=1.2)

    footer(s, 7)
    notes(s, "【3:30–4:05 性能】\n"
             "这是同一块板、同一张 720p 图、同一个 Sobel。纯软件 3108 毫秒一帧，PL 整帧传输 15.8 毫秒，196.7 倍，"
             "而且两边输出逐像素完全一致，不是用精度换速度。\n"
             "中间那 659 毫秒的逐行版本要解释一句：时间几乎全花在 Python 每行一次启动和 DMA 轮询上，不是算子在算。\n"
             "另外要说清口径：主机上 numpy 只要 40 毫秒，比板子还快，所以这组数比的是同一块嵌入式板卡上的 ARM 和 PL。")
    return s


def slide_board(prs):
    s = new(prs)
    heading(s, "07", "真机", "640×480 持续出图：切预设不黑屏，偏移每次启动标定")

    chips = [("640×480 @60", "采集卡实测输出（板卡物理上限）"),
             ("741 次写帧", "150 s 内 25 次切换，>60 ms 卡顿 0 次"),
             ("218 / 0–1 px", "本次标定值 / 残差；相位随流启动变化")]
    x = M
    for big, small in chips:
        b = box(s, x, 1.62, 3.82, 1.06, fill=CARD)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = In(0.2)
        line(tf, (big, 18, True, NAVY), first=True, spacing=1.0)
        line(tf, (small, 10, False, GREY), space_before=3, spacing=1.05)
        x += 3.82 + 0.215

    pics = [(IMG_PLAIN, "原始（pass）"), (IMG_BRIGHT, "提亮（bright）"), (IMG_SPLIT, "分屏：左原始、右增强")]
    x = 1.10
    for path, cap in pics:
        framed_picture(s, path, x, 2.92, 2.62)
        tf = frame(s, x - 0.1, 4.99, 2.82, 0.3)
        line(tf, (cap, 9.5, False, GREY), first=True, spacing=1.0, align=PP_ALIGN.CENTER)
        x += 2.62 + 0.62

    b = box(s, M, 5.44, CW, 1.10, fill=None, outline=LINE)
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.26)
    tf.margin_right = In(0.26)
    line(tf, [("①  ", 11.5, True, AMBER),
              ("HLS 把 6 个配置锁进流水线 prologue，换预设无效", 11.5, True, INK),
              (" → 配置搬进 BRAM，每像素重采样", 11.5, False, GREY)],
         first=True, spacing=1.2)
    line(tf, [("②  ", 11.5, True, AMBER),
              ("缺帧起始 SOF 时相位自由跑（历史 0 / 45 / 89）", 11.5, True, INK),
              (" → SOF 收敛到亚行内；每次启动标定（本次 218）", 11.5, False, GREY)],
         space_before=6, spacing=1.2)

    spacer(s, 6.66, "端到端（含采集与显示）未实测：这块板 HDMI 输出只到 640×480，直通不可行 —— 硬件上限，原厂位流同样如此。",
           9.5, MUTED)

    footer(s, 8)
    notes(s, "【4:05–4:35 真机】\n"
             "真机是 640×480、60 帧，采集卡实拍稳定。切预设不黑屏这件事是量化验证的：150 秒里切 25 次，写帧 741 次，"
             "超过 60 毫秒的卡顿 0 次，延迟全程恒定 21 拍。\n"
             "下面三张图是实拍，最右边那张是同一帧里左边原始、右边增强，能直接看出内核在处理像素。\n"
             "两个真问题也放在这：HLS 把 6 个配置锁进了流水线 prologue，换预设根本没生效，最后把配置搬进 BRAM；"
             "偏移看起来是随机的 0/45/89 像素，根因是缺帧起始；加 SOF 后相位只落在亚行内，但仍随流启动变化，"
             "所以启动后测一次、写补偿（本次 218），残差 0—1 像素。")
    return s


def slide_conclusion(prs):
    s = new(prs)
    heading(s, "08", "结论", "三句话收尾")

    rows = [("01", "架构", "零帧缓存流水线，延迟压到 21 个时钟；代价是放弃帧级后处理。"),
            ("02", "证据", "软硬件同算法、0 容差逐位比对；Sobel 相对 PS 快 196.7 倍。"),
            ("03", "复现", "全流程脚本化，5 个 PYNQ Skill 可迁到别的板卡与题目。")]
    y = 1.72
    for num, tag, body in rows:
        tf = frame(s, M + 0.18, y, 0.8, 0.4)
        line(tf, (num, 14, True, SKY), first=True, spacing=1.0)
        tf = frame(s, M + 0.95, y - 0.06, 1.0, 0.4)
        line(tf, (tag, 16, True, NAVY), first=True, spacing=1.0)
        tf = frame(s, M + 2.0, y - 0.06, 5.7, 0.8)
        line(tf, (body, 13.5, False, GREY), first=True, spacing=1.2)
        y += 0.95

    b = box(s, 8.55, 1.62, 4.06, 2.55, fill=CARD)
    tf = b.text_frame
    tf.margin_left = In(0.26)
    tf.margin_top = In(0.24)
    line(tf, ("交付物", 12, True, ACCENT), first=True, spacing=1.0)
    for item in ["enhance_hdmi.bit（WNS +0.206 ns）",
                 "源码 + 两条通路位流，两条命令重建",
                 "设计报告 · 一页海报 · 5 个 Skill",
                 "DMA 逐位比对脚本与原始 metrics"]:
        line(tf, ("· " + item, 10, False, GREY), space_before=9, spacing=1.2)

    b = box(s, M, 4.86, CW, 0.88, fill=None, outline=LINE)
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = In(0.26)
    line(tf, [("已知限制  ", 11, True, AMBER),
              ("HDMI 只能 640×480（板卡上限）· 端到端与功耗未实测 · Sobel/高斯/锐化三算子并行未集成进同一 overlay",
               11, False, GREY)], first=True, spacing=1.0)

    tf = frame(s, M + 0.18, 6.02, 7.0, 0.65)
    line(tf, ("谢谢，欢迎提问", 26, True, NAVY), first=True, spacing=1.0)

    footer(s, 9)
    notes(s, "【4:35–5:00 收尾】\n"
             "三句话：架构上零帧缓存，延迟 21 个时钟，代价是放弃帧级后处理；证据上软硬件同算法、0 容差比对，"
             "相对 ARM 196.7 倍；工程上全流程脚本化，五个 Skill 可以迁移。\n"
             "限制也如实说：这块板的 HDMI 只能 640×480，端到端和功耗还没测，三个算子还没进同一个 overlay。\n"
             "谢谢各位老师，欢迎提问。")
    return s


def normalise_slide_type(path):
    """python-pptx keeps the template's legacy type="screen4x3" next to the 16:9 size."""
    import re
    import zipfile

    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/presentation.xml":
                text = data.decode("utf-8")
                text = re.sub(r'<p:sldSz([^>]*?)type="[^"]*"',
                              r'<p:sldSz\1type="screen16x9"', text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


def main() -> int:
    prs = Presentation()
    prs.slide_width = Emu(12192000)     # exactly PowerPoint's 16:9 preset
    prs.slide_height = Emu(6858000)

    slide_cover(prs)
    slide_problem(prs)
    slide_arch(prs)
    slide_kernel(prs)
    slide_effect(prs)
    slide_correctness(prs)
    slide_perf(prs)
    slide_board(prs)
    slide_conclusion(prs)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    normalise_slide_type(OUT)
    print(f"wrote {OUT}  ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
